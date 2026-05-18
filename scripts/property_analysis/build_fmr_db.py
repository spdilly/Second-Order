"""
Build the HUD FMR SQLite database from the official xlsx files.

Reads:
  templates/property_analysis/data/FY26_FMRs.xlsx       (county/metro-level, 4,764 rows)
  templates/property_analysis/data/FY2026_SAFMRs.xlsx   (ZIP-level, 51,895 rows)

Writes:
  templates/property_analysis/data/hud_fmr.db
    - safmr        ZIP-level rents (primary lookup for SAFMR metros)
    - fmr          County/metro-level rents (fallback for non-SAFMR areas)
    - metadata     Build timestamp + source file fingerprints

Run annually after downloading the new FY xlsx files from huduser.gov.

Usage:
  python -m scripts.property_analysis.build_fmr_db
"""
from __future__ import annotations

import hashlib
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = ROOT / "templates" / "property_analysis" / "data"
DB_PATH = DATA_DIR / "hud_fmr.db"
FMR_XLSX = DATA_DIR / "FY26_FMRs.xlsx"
SAFMR_XLSX = DATA_DIR / "FY2026_SAFMRs.xlsx"


def file_fingerprint(path: Path) -> str:
    h = hashlib.sha256()
    h.update(path.read_bytes())
    return h.hexdigest()[:16]


def build_schema(conn: sqlite3.Connection) -> None:
    conn.executescript("""
        DROP TABLE IF EXISTS safmr;
        DROP TABLE IF EXISTS fmr;
        DROP TABLE IF EXISTS metadata;

        CREATE TABLE safmr (
            zip            TEXT NOT NULL,
            hud_area_code  TEXT NOT NULL,
            hud_area_name  TEXT NOT NULL,
            fmr_0br        INTEGER,
            fmr_1br        INTEGER,
            fmr_2br        INTEGER,
            fmr_3br        INTEGER,
            fmr_4br        INTEGER,
            PRIMARY KEY (zip, hud_area_code)
        );
        CREATE INDEX idx_safmr_zip ON safmr(zip);

        CREATE TABLE fmr (
            stusps         TEXT NOT NULL,
            state_fips     TEXT NOT NULL,
            county_fips    TEXT NOT NULL,
            full_fips      TEXT NOT NULL,
            hud_area_code  TEXT NOT NULL,
            countyname     TEXT NOT NULL,
            hud_area_name  TEXT NOT NULL,
            is_metro       INTEGER,
            population     INTEGER,
            fmr_0br        INTEGER,
            fmr_1br        INTEGER,
            fmr_2br        INTEGER,
            fmr_3br        INTEGER,
            fmr_4br        INTEGER,
            PRIMARY KEY (full_fips)
        );
        CREATE INDEX idx_fmr_state_county ON fmr(stusps, countyname);
        CREATE INDEX idx_fmr_hud_area ON fmr(hud_area_code);

        CREATE TABLE metadata (
            key   TEXT PRIMARY KEY,
            value TEXT
        );
    """)


def load_safmr(conn: sqlite3.Connection) -> int:
    wb = openpyxl.load_workbook(SAFMR_XLSX, read_only=True, data_only=True)
    ws = wb["SAFMRs"]
    rows_inserted = 0
    rows = ws.iter_rows(min_row=2, values_only=True)
    batch: list[tuple] = []
    for row in rows:
        zip_code = str(row[0]).strip() if row[0] is not None else None
        if not zip_code:
            continue
        zip_code = zip_code.zfill(5)
        batch.append((
            zip_code,
            row[1],          # HUD Area Code
            row[2],          # HUD Area Name
            row[3],          # SAFMR 0BR
            row[6],          # SAFMR 1BR
            row[9],          # SAFMR 2BR
            row[12],         # SAFMR 3BR
            row[15],         # SAFMR 4BR
        ))
        if len(batch) >= 1000:
            conn.executemany(
                "INSERT OR REPLACE INTO safmr VALUES (?,?,?,?,?,?,?,?)",
                batch,
            )
            rows_inserted += len(batch)
            batch.clear()
    if batch:
        conn.executemany(
            "INSERT OR REPLACE INTO safmr VALUES (?,?,?,?,?,?,?,?)",
            batch,
        )
        rows_inserted += len(batch)
    wb.close()
    return rows_inserted


def load_fmr(conn: sqlite3.Connection) -> int:
    wb = openpyxl.load_workbook(FMR_XLSX, read_only=True, data_only=True)
    ws = wb["FY26_FMRs"]
    rows_inserted = 0
    rows = ws.iter_rows(min_row=2, values_only=True)
    batch: list[tuple] = []
    for row in rows:
        if row[0] is None:
            continue
        state_fips = str(row[1]).zfill(2)
        fips_str = str(row[7])
        county_fips = fips_str[2:5] if len(fips_str) >= 5 else ""
        batch.append((
            row[0],                  # stusps
            state_fips,              # state_fips
            county_fips,             # county_fips
            fips_str,                # full_fips
            row[2],                  # hud_area_code
            row[3],                  # countyname
            row[6],                  # hud_area_name
            int(row[5]) if row[5] is not None else None,  # is_metro
            int(row[8]) if row[8] is not None else None,  # population
            row[9], row[10], row[11], row[12], row[13],
        ))
        if len(batch) >= 1000:
            conn.executemany(
                "INSERT OR REPLACE INTO fmr VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                batch,
            )
            rows_inserted += len(batch)
            batch.clear()
    if batch:
        conn.executemany(
            "INSERT OR REPLACE INTO fmr VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            batch,
        )
        rows_inserted += len(batch)
    wb.close()
    return rows_inserted


def write_metadata(conn: sqlite3.Connection, safmr_n: int, fmr_n: int) -> None:
    rows = [
        ("built_at", datetime.utcnow().isoformat(timespec="seconds") + "Z"),
        ("fy", "2026"),
        ("safmr_rows", str(safmr_n)),
        ("fmr_rows", str(fmr_n)),
        ("safmr_xlsx_sha256_16", file_fingerprint(SAFMR_XLSX)),
        ("fmr_xlsx_sha256_16", file_fingerprint(FMR_XLSX)),
    ]
    conn.executemany("INSERT INTO metadata VALUES (?, ?)", rows)


def build() -> None:
    if not FMR_XLSX.exists():
        print(f"ERROR: missing {FMR_XLSX}")
        sys.exit(1)
    if not SAFMR_XLSX.exists():
        print(f"ERROR: missing {SAFMR_XLSX}")
        sys.exit(1)

    if DB_PATH.exists():
        DB_PATH.unlink()
    conn = sqlite3.connect(DB_PATH)
    try:
        build_schema(conn)
        print("Loading SAFMR (ZIP-level)...")
        safmr_n = load_safmr(conn)
        print(f"  {safmr_n:,} rows")
        print("Loading FMR (county-level)...")
        fmr_n = load_fmr(conn)
        print(f"  {fmr_n:,} rows")
        write_metadata(conn, safmr_n, fmr_n)
        conn.commit()
    finally:
        conn.close()
    print(f"\nBuilt {DB_PATH} ({DB_PATH.stat().st_size/1024/1024:.2f} MB)")


if __name__ == "__main__":
    build()
