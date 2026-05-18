"""
Populate the SFH proforma template with a specific property's inputs.

Reads the empty template, writes inputs to the named ranges on the Inputs tab,
and saves to a per-property output filename. Never touches calculation cells —
all model logic is in the template's formulas, which recalculate on Excel open.

Usage:
    from scripts.property_analysis.populate import populate_proforma

    output_path = populate_proforma(
        inputs=Inputs(...),
        output_dir="output/projects/Joe Berlin",
        property_label="1417_S_Canal_St",
    )
"""
from __future__ import annotations

import shutil
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from scripts.property_analysis.compute import Inputs

PROJECT_ROOT = Path(__file__).resolve().parents[2]
TEMPLATE_PATH = PROJECT_ROOT / "templates" / "property_analysis" / "proforma_template.xlsx"


def populate_proforma(
    inputs: Inputs,
    output_dir: Path | str,
    property_label: str,
    template_path: Path | None = None,
) -> Path:
    """
    Copy the proforma template to a per-property output file and write inputs.

    The xlsx is written INTO output_dir directly. The caller (analyze.py)
    decides whether output_dir is a per-deal packet folder or a flat path —
    populate_proforma does not impose folder structure beyond ensuring the
    directory exists.

    Args:
        inputs: Fully-specified Inputs dataclass.
        output_dir: Where to write the populated xlsx.
        property_label: Slugified label for the filename (e.g. "1417_S_Canal_St").
        template_path: Override the template location.

    Returns:
        Path to the populated xlsx.
    """
    template_path = template_path or TEMPLATE_PATH
    if not template_path.exists():
        raise FileNotFoundError(
            f"Proforma template not found at {template_path}. "
            "Run: python -m financial_model_framework.templates.sfh_proforma.builder"
        )

    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    output_path = output_dir / f"{property_label}_proforma.xlsx"
    # If a packet folder already has this file (rare — folder name is per-run
    # timestamp), still avoid overwriting by appending HHMMSS.
    if output_path.exists():
        stamp = datetime.now().strftime("%H%M%S")
        output_path = output_dir / f"{property_label}_proforma_{stamp}.xlsx"
        counter = 1
        while output_path.exists():
            output_path = output_dir / f"{property_label}_proforma_{stamp}_{counter}.xlsx"
            counter += 1
    shutil.copyfile(template_path, output_path)

    wb = load_workbook(output_path)
    # Build map: named range -> (sheet, row, col) from the workbook itself
    name_to_addr: dict[str, tuple[str, int, int]] = {}
    for name in wb.defined_names:
        defn = wb.defined_names[name]
        # attr_text is like 'Inputs'!$C$12
        attr = defn.attr_text
        if not attr:
            continue
        try:
            sheet_part, cell_part = attr.split("!", 1)
            sheet_name = sheet_part.strip("'")
            # Strip $ from cell ref like $C$12
            ref = cell_part.replace("$", "").split(":")[0]
            from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
            col_letter, row_num = coordinate_from_string(ref)
            col_idx = column_index_from_string(col_letter)
            name_to_addr[name] = (sheet_name, row_num, col_idx)
        except Exception:
            continue

    # Write each input field value to its named range cell
    input_dict = asdict(inputs)
    written = 0
    skipped = []
    for field_name, value in input_dict.items():
        if field_name not in name_to_addr:
            # Some inputs (address, zip) may not have named ranges if we
            # decide to keep them informational only — but they should
            skipped.append(field_name)
            continue
        sheet, row, col = name_to_addr[field_name]
        wb[sheet].cell(row=row, column=col, value=value)
        written += 1

    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    # Smoke test on 1417 S Canal St using Joe's existing assumptions
    inputs = Inputs(
        address="1417 S Canal St, Pittsburgh, PA 15215",
        zip="15215",
        beds=3,
        purchase_price=130000,
        arv=255000,
        rehab_cost=55000,
        closing_cost_pct=0.07,
        holding_months=3,
        down_payment_pct=1.0,
        interest_rate=0.095,
        loan_term_yrs=30,
        refi_year=3,
        refi_ltv=0.75,
        refi_interest_rate=0.075,
        refi_term_yrs=30,
        refi_closing_cost_pct=0.04,
        monthly_rent=2500,
        rent_growth=0.04,
        vacancy_rate=0.04,
        management_pct=0.10,
        maintenance_pct=0.04,
        capex_reserve_pct=0.0,
        property_tax_annual=3249,
        insurance_annual=1020,
        hoa_annual=0,
        utilities_annual=0,
        expense_growth=0.03,
        hold_period_yrs=10,
        exit_cap_rate=0.08,
        cost_of_sale_pct=0.01,
    )
    out = populate_proforma(
        inputs=inputs,
        output_dir=PROJECT_ROOT / "output" / "projects" / "Joe Berlin",
        property_label="1417_S_Canal_St_TEST",
    )
    print(f"Populated: {out}")
    print(f"Size: {out.stat().st_size/1024:.1f} KB")
